"""Convert docs/Presentation_Script.md into docs/Presentation_Script.docx.

Run from the repository root:
    pip install python-docx
    python docs/build_script_docx.py
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DOCS = Path(__file__).resolve().parent
SRC, OUT = DOCS / "Presentation_Script.md", DOCS / "Presentation_Script.docx"

INK = RGBColor(0x1D, 0x21, 0x25)
BLUE = RGBColor(0x2F, 0x5F, 0x9E)
RUST = RGBColor(0xB5, 0x54, 0x1C)
MUTED = RGBColor(0x5F, 0x63, 0x68)


def _shade(element, hex_fill):
    pr = element.get_or_add_pPr() if hasattr(element, "get_or_add_pPr") else element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    pr.append(shd)


def inline(paragraph, line, color=None, italic=False, size=None):
    """Add text where **bold**, *italic* and `code` spans are formatted."""
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", line):
        if not part:
            continue
        if part.startswith("**"):
            r = paragraph.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("*"):
            r = paragraph.add_run(part[1:-1])
            r.italic = True
            r.font.color.rgb = MUTED
        elif part.startswith("`"):
            r = paragraph.add_run(part[1:-1])
            r.font.name = "Consolas"
        else:
            r = paragraph.add_run(part)
        if italic:
            r.italic = True
        if color is not None and not part.startswith("*"):
            r.font.color.rgb = color
        if size:
            r.font.size = Pt(size)


def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{3,}:?", c) for c in r)]
    t = doc.add_table(rows=len(cells), cols=len(cells[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(cells):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            p = cell.paragraphs[0]
            inline(p, val, size=9.5)
            if i == 0:
                for r in p.runs:
                    r.bold = True
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                _shade(cell._tc, "1F4E82")
            elif i % 2 == 0:
                _shade(cell._tc, "F2F5FA")
    doc.add_paragraph()


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, color in (("Title", 24, INK), ("Heading 1", 16, BLUE), ("Heading 2", 15, BLUE),
                              ("Heading 3", 12.5, RUST)):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = True

    lines = SRC.read_text(encoding="utf-8").splitlines()
    para, table, first_slide = [], [], True

    def flush():
        if para:
            text = " ".join(s.strip() for s in para)
            para.clear()
            if text.startswith("**Speaker:**"):
                p = doc.add_paragraph()
                _shade(p._p, "E4EDF7")
                inline(p, text, color=BLUE)
            elif text.startswith("**Handover:**"):
                p = doc.add_paragraph()
                inline(p, text, color=RUST, italic=True)
            elif text.startswith("*[") and text.endswith("]*") and text.count("*[") == 1:
                p = doc.add_paragraph()
                inline(p, text)
            else:
                inline(doc.add_paragraph(), text)
        if table:
            add_table(doc, table)
            table.clear()

    for raw in lines:
        line = raw.rstrip()
        if line.startswith("|"):
            if para:
                flush()
            table.append(line)
            continue
        if table:
            flush()
        if not line.strip() or line.strip() == "---":
            flush()
        elif line.startswith("# "):
            flush()
            doc.add_paragraph(line[2:], style="Title")
        elif line.startswith("## "):
            flush()
            title = line[3:]
            if title.startswith("Slide "):
                if not first_slide:
                    doc.add_paragraph()
                first_slide = False
            doc.add_heading(title, level=2).paragraph_format.keep_with_next = True
        elif line.startswith("### "):
            flush()
            doc.add_heading(line[4:], level=3).paragraph_format.keep_with_next = True
        elif line.startswith("- "):
            flush()
            inline(doc.add_paragraph(style="List Bullet"), line[2:])
        elif line.startswith("  ") and doc.paragraphs and doc.paragraphs[-1].style.name == "List Bullet":
            inline(doc.paragraphs[-1], " " + line.strip())
        else:
            para.append(line)
    flush()

    footer = sec.footer.paragraphs[0]
    footer.text = "Term Deposit Predictor · Data Miners · Group 1.1 · Presentation script"
    footer.runs[0].font.size = Pt(9)
    footer.runs[0].font.color.rgb = MUTED
    doc.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    build()
